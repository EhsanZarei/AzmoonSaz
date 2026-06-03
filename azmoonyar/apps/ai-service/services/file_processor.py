import logging
from typing import Optional

logger = logging.getLogger(__name__)


async def process_file(content: bytes, content_type: str, filename: Optional[str] = None) -> str:
    """استخراج متن از فایل‌های مختلف"""

    if content_type == "application/pdf":
        return await extract_from_pdf(content)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return await extract_from_docx(content)
    elif content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return await extract_from_pptx(content)
    elif content_type.startswith("image/"):
        return await extract_from_image(content)
    else:
        raise ValueError(f"فرمت پشتیبانی نمی‌شود: {content_type}")


async def extract_from_pdf(content: bytes) -> str:
    """استخراج متن از PDF"""
    try:
        import pdfplumber
        import io

        text_parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        return "\n\n".join(text_parts)
    except ImportError:
        logger.warning("pdfplumber not installed, trying pypdf2")
        try:
            import PyPDF2
            import io

            reader = PyPDF2.PdfReader(io.BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            raise Exception("کتابخانه PDF نصب نشده است")


async def extract_from_docx(content: bytes) -> str:
    """استخراج متن از Word"""
    try:
        from docx import Document
        import io

        doc = Document(io.BytesIO(content))
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    except ImportError:
        raise Exception("کتابخانه python-docx نصب نشده است")


async def extract_from_pptx(content: bytes) -> str:
    """استخراج متن از PowerPoint"""
    try:
        from pptx import Presentation
        import io

        prs = Presentation(io.BytesIO(content))
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text)

        return "\n".join(text_parts)
    except ImportError:
        raise Exception("کتابخانه python-pptx نصب نشده است")


async def extract_from_image(content: bytes) -> str:
    """استخراج متن از تصویر با OCR"""
    try:
        import pytesseract
        from PIL import Image
        import io

        image = Image.open(io.BytesIO(content))
        # OCR با پشتیبانی فارسی
        text = pytesseract.image_to_string(image, lang="fas+eng")
        return text.strip()
    except ImportError:
        raise Exception("کتابخانه pytesseract نصب نشده است")
    except Exception as e:
        raise Exception(f"خطا در OCR: {str(e)}")
